"""Phase 42: a date reads dd-mm-yyyy wherever it is read, shown or saved."""

import datetime
import io

import pandas as pd
from openpyxl import Workbook, load_workbook

from cleaner.export import build_workbook
from cleaner.loaders import read_table
from cleaner.profiling import detect_column_type, parse_datetime_series
from dashboard.html_export import frame_to_html
from utils.dates import (
    DATE_DISPLAY_FORMAT,
    DATETIME_DISPLAY_FORMAT,
    date_column_config,
    dates_as_text,
    excel_ready,
    format_date_value,
)


def _workbook_bytes(rows: list[list]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


class TestReadingExcel:
    def test_a_date_cell_reads_as_dd_mm_yyyy_not_pandas_own_form(self):
        data = _workbook_bytes([["on"], [datetime.datetime(2025, 4, 3)]])
        assert read_table(data, "book.xlsx")["on"].iloc[0] == "03-04-2025"

    def test_a_time_is_kept_only_when_there_is_one(self):
        data = _workbook_bytes([["on"], [datetime.datetime(2025, 4, 3, 10, 30)]])
        assert read_table(data, "book.xlsx")["on"].iloc[0] == "03-04-2025 10:30:00"

    def test_everything_else_is_text_exactly_as_before(self):
        data = _workbook_bytes([["code", "qty"], ["007", 1.0], [None, 2.5], ["x", None]])
        frame = read_table(data, "book.xlsx")
        assert frame["code"].iloc[0] == "007"
        assert frame["qty"].iloc[0] == "1"
        assert frame["qty"].iloc[1] == "2.5"
        assert pd.isna(frame["code"].iloc[1])
        assert pd.isna(frame["qty"].iloc[2])


class TestReadingADateWithNoStatedFormat:
    def test_day_comes_first(self):
        parsed, failed = parse_datetime_series(pd.Series(["03-04-2025", "13-04-2025"]))
        assert list(parsed) == [pd.Timestamp("2025-04-03"), pd.Timestamp("2025-04-13")]
        assert len(failed) == 0

    def test_iso_dates_still_read_as_iso(self):
        parsed, _ = parse_datetime_series(pd.Series(["2025-09-03", "2025-12-31"]))
        assert parsed.iloc[0] == pd.Timestamp("2025-09-03")

    def test_a_day_first_column_is_detected_as_dates(self):
        assert detect_column_type(pd.Series(["03-04-2025", "13-04-2025", "31-12-2024"] * 5)) == "date"


class TestShowingOnScreen:
    def test_a_date_only_column_is_shown_dd_mm_yyyy(self):
        frame = pd.DataFrame({"on": pd.to_datetime(["2025-04-03"]), "n": [1]})
        config = date_column_config(frame)
        assert list(config) == ["on"]
        assert config["on"]["type_config"]["format"] == DATE_DISPLAY_FORMAT

    def test_a_column_with_times_keeps_the_time(self):
        frame = pd.DataFrame({"on": pd.to_datetime(["2025-04-03 10:30"])})
        assert date_column_config(frame)["on"]["type_config"]["format"] == DATETIME_DISPLAY_FORMAT

    def test_a_column_the_page_already_configured_is_left_alone(self):
        frame = pd.DataFrame({"on": pd.to_datetime(["2025-04-03"])})
        assert date_column_config(frame, {"on": "Mine"}) == {"on": "Mine"}


class TestSaving:
    def test_one_date_as_text(self):
        assert format_date_value(datetime.date(2025, 4, 3)) == "03-04-2025"

    def test_csv_dates_are_written_dd_mm_yyyy(self):
        frame = pd.DataFrame({"on": pd.to_datetime(["2025-04-03", None]), "n": [1, 2]})
        assert dates_as_text(frame).to_csv(index=False).splitlines() == ["on,n", "03-04-2025,1", ",2"]

    def test_excel_keeps_real_dates_formatted_dd_mm_yyyy(self):
        frame = pd.DataFrame(
            {"on": pd.to_datetime(["2025-04-03"]), "at": pd.to_datetime(["2025-04-03 10:30"])}
        )
        sheet = load_workbook(io.BytesIO(build_workbook([("Data", frame)]))).worksheets[0]
        assert sheet["A2"].value == datetime.datetime(2025, 4, 3)
        assert sheet["A2"].number_format == "dd-mm-yyyy"
        assert sheet["B2"].number_format == "dd-mm-yyyy hh:mm:ss"

    def test_excel_ready_drops_a_time_zone_excel_would_refuse(self):
        frame = pd.DataFrame({"on": pd.to_datetime(["2025-04-03 10:30"]).tz_localize("Asia/Kolkata")})
        assert excel_ready(frame)["on"].dt.tz is None

    def test_the_html_report_table_writes_dd_mm_yyyy(self):
        html = frame_to_html(pd.DataFrame({"on": pd.to_datetime(["2025-04-03"])}))
        assert "03-04-2025" in html
